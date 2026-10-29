"""實驗一抽樣（手冊步驟 5、第 8.1 節）：給標註員判斷「兩則素材是不是同一事件」的 100 組配對。

流程（全部離線、不呼叫 LLM、來源 DB 只用唯讀 URI 開）：

1. 從來源 DB 讀素材；每則素材歸到「出刊日」：`first_seen_at` 落在 (D-1 06:00, D 06:00] 台北 → D。
   dev＝10/27–10/30 出刊日、test＝10/31–11/02 出刊日；兩則分屬 dev、test 的配對不抽。
2. 兩種方法各在暫存 SQLite 重跑一次增量聚合（每 `round_hours` 一輪，只放入該輪之前收到的素材）：
   - 方法 A（現行版）：`cluster_recent_items` 用 `METHOD_A_RULES`（相似度 ≥ 0.28 就合併、不需共同實體＝舊規則），
     重複稿用收集當時的標記，已停用來源也參與（當時還沒停用）；跨語言用來源 DB 各期
     `daily_editions.stats.merged_groups`（find_same_events 的結果）重播：live cluster id → 素材 → A 的事件。
   - 方法 C（步驟 3 現行版）：設定檔的兩段門檻＋共同實體＋譯名表、重算重複稿、排除停用來源；
     每個出刊日 06:00 用 `desk_merges`（譯名表＋共同實體的跨語言規則，不含 LLM 第二意見）合併。
   「系統合併」＝兩則素材落在同一事件，含遞移（經第三則或跨語言合併間接相連）。
3. 候選配對：兩則都可標（啟用來源、非重複稿、中文或英文、有標題、同一 split）、發布時間差 ≤ 48 小時。
   分層只看程式判斷：A、C 都合併／只有 A／只有 C／都不合併但 TF-IDF 相似度 ≥ 0.15。
   同語言 4 層各 15、中英 4 層各 10（dev、test 各半，見 `QUOTAS`）。中英某層不夠時，用 48 小時內、
   同版別的隨機中英配對補（`filled=True`）；同語言不足就照實少抽。
4. dev 先放錯誤案例（`tests/regression/cases`）相關的配對（每案最多 1 組），test 不放任何錯誤案例素材。
5. B 複標：dev、test 各隨機 15 組。所有隨機都來自同一個 `random.Random(seed)`，同樣輸入＝同樣輸出。
"""

from __future__ import annotations

import csv
import hashlib
import html
import json
import random
import re
import sqlite3
import tempfile
import unicodedata
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import zhconv
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from ..collection.clustering import _cosine, _document_text, _vectorize, cluster_recent_items, lang_group, tokenize
from ..collection.collector import recompute_duplicates
from ..collection.entities import EntityLexicon, MergeRules, load_lexicon
from ..models import Base, SourceItem
from ..seeds.news_sources import SourceRegistry

TAIPEI = ZoneInfo("Asia/Taipei")
# 固定亂數種子（手冊步驟 5：程式與種子先 commit 再抽）
SEED = 20261104
EDITION_CUTOFF = time(6, 0)
DEV_DATES = (date(2026, 10, 27), date(2026, 10, 30))
TEST_DATES = (date(2026, 10, 31), date(2026, 11, 2))
MAX_GAP = timedelta(hours=48)
MIN_SIMILARITY = 0.15
SUMMARY_CHARS = 300
STRATA = ("both", "only_a", "only_c", "neither")
STRATUM_LABELS = {"both": "A、C 都合併", "only_a": "只有 A 合併", "only_c": "只有 C 合併", "neither": "都不合併、相似度≥0.15"}
# 每層 dev／test 名額：同語言每層 15（8＋7 交錯）、中英每層 10（5＋5）→ dev 50、test 50
QUOTAS = {
    "dev": {"same": {"both": 8, "only_a": 7, "only_c": 8, "neither": 7}, "cross": dict.fromkeys(STRATA, 5)},
    "test": {"same": {"both": 7, "only_a": 8, "only_c": 7, "neither": 8}, "cross": dict.fromkeys(STRATA, 5)},
}
B_PER_SPLIT = 15
# 方法 A：舊規則（門檻 0.28、不看共同實體）＝現行聚合程式的參數設定，不另抄一份
METHOD_A_RULES = MergeRules(high_threshold=0.28, low_threshold=0.28, high_min_entities=0, mid_min_entities=0)
ANNOTATION_COLUMNS = ("pair_id", "標題 1", "摘要 1", "時間 1", "標題 2", "摘要 2", "時間 2", "判斷", "一句理由")
JUDGMENTS = ("是", "不是", "不確定")
LANGS = ("zh", "en")


@dataclass(frozen=True)
class Item:
    id: int
    source_id: str
    section: str
    tier: str
    title: str
    summary: str
    lang: str
    url: str
    canonical_url: str
    published_at: datetime  # UTC aware
    published_at_estimated: bool
    first_seen_at: datetime  # UTC aware
    duplicate_of_id: int | None
    cluster_id: int | None  # 來源 DB 的（live）事件 id，給方法 A 重播 merged_groups 用

    @property
    def group(self) -> str:
        return lang_group(self.lang)


@dataclass
class Pair:
    a: int
    b: int
    split: str
    kind: str  # same／cross
    a_merged: bool
    c_merged: bool
    similarity: float
    stratum: str | None = None
    filled: bool = False
    error_case: str | None = None

    @property
    def key(self) -> tuple[int, int]:
        return (self.a, self.b)


@dataclass
class SampleResult:
    selected: list[Pair]
    b_pairs: list[Pair]
    pair_ids: dict[tuple[int, int], str]
    shortages: list[dict]
    pools: dict[str, int]
    merge_replay: dict[str, str]
    notes: list[str] = field(default_factory=list)


# ── 讀資料 ───────────────────────────────────────────────────────────────


def _utc(text: str | None) -> datetime | None:
    if not text:
        return None
    value = datetime.fromisoformat(text)
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def readonly_uri(path: str | Path) -> str:
    """母本是唯讀檔：`mode=ro` 不夠，要加 `immutable=1` 才開得起來。"""
    return f"file:{Path(path).resolve()}?mode=ro&immutable=1"


def load_items(db_path: str | Path, *, until: datetime) -> list[Item]:
    conn = sqlite3.connect(readonly_uri(db_path), uri=True)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            "SELECT id, source_id, section, tier, title, summary, lang, url, canonical_url, published_at,"
            " published_at_estimated, first_seen_at, duplicate_of_id, cluster_id FROM source_items"
            " WHERE first_seen_at <= ? ORDER BY first_seen_at, id",
            (until.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S.%f"),),
        ).fetchall()
    finally:
        conn.close()
    return [
        Item(
            r["id"], r["source_id"], r["section"], r["tier"], r["title"] or "", r["summary"] or "", r["lang"] or "",
            r["url"], r["canonical_url"], _utc(r["published_at"]), bool(r["published_at_estimated"]),
            _utc(r["first_seen_at"]), r["duplicate_of_id"], r["cluster_id"],
        )
        for r in rows
    ]


def load_merged_groups(db_path: str | Path, dates: Iterable[date]) -> dict[date, list[list[int]] | None]:
    """各出刊日 `stats.merged_groups`（live cluster id）；沒有那期或沒有紀錄是 None。"""
    conn = sqlite3.connect(readonly_uri(db_path), uri=True)
    try:
        rows = dict(conn.execute("SELECT edition_date, stats FROM daily_editions").fetchall())
    finally:
        conn.close()
    out = {}
    for day in dates:
        stats = json.loads(rows[day.isoformat()]) if rows.get(day.isoformat()) else None
        out[day] = (stats or {}).get("merged_groups") if stats else None
    return out


def edition_of(first_seen_at: datetime) -> date:
    """`first_seen_at` 在 (D-1 06:00, D 06:00] 台北 → 出刊日 D（那天 06:00 準備的報紙會用到它）。"""
    local = first_seen_at.astimezone(TAIPEI)
    return local.date() if local.time() <= EDITION_CUTOFF else local.date() + timedelta(days=1)


def edition_cutoff(day: date) -> datetime:
    return datetime.combine(day, EDITION_CUTOFF, TAIPEI).astimezone(timezone.utc)


def date_range(first: date, last: date) -> list[date]:
    return [first + timedelta(days=i) for i in range((last - first).days + 1)]


def split_of(day: date, dev: tuple[date, date] = DEV_DATES, test: tuple[date, date] = TEST_DATES) -> str | None:
    if dev[0] <= day <= dev[1]:
        return "dev"
    if test[0] <= day <= test[1]:
        return "test"
    return None


# ── 兩種方法的聚合 ───────────────────────────────────────────────────────


def _row(item: Item) -> SourceItem:
    return SourceItem(
        id=item.id, source_id=item.source_id, section=item.section, tier=item.tier, item_key=str(item.id),
        url=item.url, canonical_url=item.canonical_url, title=item.title, summary=item.summary, lang=item.lang,
        published_at=item.published_at, published_at_estimated=item.published_at_estimated,
        first_seen_at=item.first_seen_at, duplicate_of_id=item.duplicate_of_id,
    )


def round_times(items: list[Item], end: datetime, round_hours: float, extra: Iterable[datetime] = ()) -> list[datetime]:
    """從第一則素材起每 `round_hours` 一輪，加上 `extra`（出刊日 06:00），最後一輪是 `end`。"""
    start = min(i.first_seen_at for i in items)
    times, t = set(), start
    while t < end:
        times.add(t)
        t += timedelta(hours=round_hours)
    times.add(end)
    times |= {x for x in extra if start <= x <= end}
    return sorted(times)


def simulate(
    items: list[Item],
    *,
    rules: MergeRules,
    lexicon: EntityLexicon,
    registry: SourceRegistry,
    recompute: bool,
    disabled_sources: frozenset[str],
    end: datetime,
    round_hours: float,
    edition_dates: Iterable[date] = (),
    desk: bool = False,
) -> tuple[dict[int, int | None], list[list[int]]]:
    """在暫存 SQLite 照收到的順序逐輪聚合。回傳（素材 → 事件 id, 出刊日 desk 合併組）。"""
    editions = {edition_cutoff(d): d for d in edition_dates}
    rows = [_row(i) for i in items]
    if recompute:
        recompute_duplicates(rows)
    desk_groups: list[list[int]] = []
    with tempfile.TemporaryDirectory(prefix="exp1-") as tmp:
        engine = create_engine(f"sqlite:///{Path(tmp) / 'exp1.db'}", future=True)
        Base.metadata.create_all(engine)
        pending = sorted(rows, key=lambda r: (r.first_seen_at, r.id))
        with Session(engine, autoflush=False) as session:
            for now in round_times(items, end, round_hours, editions):
                batch = [r for r in pending if r.first_seen_at <= now]
                pending = pending[len(batch):]
                session.add_all(batch)
                session.flush()
                cluster_recent_items(
                    session, now=now, rules=rules, lexicon=lexicon,
                    general_sources=registry.general_source_ids(), disabled_sources=disabled_sources,
                )
                if desk and now in editions:
                    desk_groups += _desk_groups(session, editions[now], now, rules, lexicon, registry, disabled_sources)
                session.commit()
            assert not pending
            cluster_of = {r.id: r.cluster_id for r in rows}
        engine.dispose()
    return cluster_of, desk_groups


def _desk_groups(session, edition_date, now, rules, lexicon, registry, disabled_sources) -> list[list[int]]:
    """方法 C 的跨語言：照線上編輯會議（評分→改寫分版→desk_pool→譯名表＋共同實體），不含 LLM 第二意見。"""
    from ..daily.desk import desk_merges
    from ..daily.scoring import desk_pool, score_candidates
    from ..daily.sections import rewrite_sections

    scored = score_candidates(
        session, edition_date=edition_date, now=now,
        general_sources=registry.general_source_ids(), disabled_sources=disabled_sources,
    )
    scored = rewrite_sections(session, scored, registry=registry)
    result = desk_merges(
        session, desk_pool(scored), [], lexicon=lexicon, rules=rules, disabled_sources=disabled_sources,
        context_cluster_ids=[s.cluster_id for s in scored],
    )
    return result.groups


class Components:
    """素材 → 系統事件（含遞移）。同一事件、或事件之間被合併組連起來，都算同一個。"""

    def __init__(self, cluster_of: dict[int, int | None], groups: Iterable[Iterable[int]]) -> None:
        self.parent: dict[int, int] = {}
        for cluster_ids in groups:
            ids = list(cluster_ids)
            for other in ids[1:]:
                self._union(ids[0], other)
        self.cluster_of = cluster_of

    def _find(self, x: int) -> int:
        self.parent.setdefault(x, x)
        while self.parent[x] != x:
            self.parent[x] = self.parent[self.parent[x]]
            x = self.parent[x]
        return x

    def _union(self, a: int, b: int) -> None:
        ra, rb = self._find(a), self._find(b)
        if ra != rb:
            self.parent[max(ra, rb)] = min(ra, rb)

    def root(self, item_id: int) -> tuple[str, int]:
        cluster = self.cluster_of.get(item_id)
        return ("item", item_id) if cluster is None else ("cluster", self._find(cluster))

    def merged(self, a: int, b: int) -> bool:
        return self.root(a) == self.root(b)


def replay_groups(merged_groups: Iterable[list[int]], live_cluster_of_items: dict[int, list[int]],
                  cluster_of: dict[int, int | None]) -> list[list[int]]:
    """live cluster id 組 → 素材 → 方法 A 的事件 id 組（同 tests/regression 的 1027_merges 重播）。"""
    groups = []
    for group in merged_groups:
        ids = list(dict.fromkeys(
            cluster_of[i] for cid in group for i in live_cluster_of_items.get(cid, []) if cluster_of.get(i) is not None
        ))
        if len(ids) >= 2:
            groups.append(ids)
    return groups


# ── 候選配對與分層 ───────────────────────────────────────────────────────


def eligible_items(items: list[Item], disabled_sources: frozenset[str], c_duplicates: dict[int, int | None],
                   dev: tuple[date, date], test: tuple[date, date]) -> dict[int, str]:
    """可以拿去給標註員看的素材 → split：啟用來源、不是重複稿、中文或英文、有標題、在 dev／test 出刊日內。"""
    out = {}
    for item in items:
        if item.source_id in disabled_sources or c_duplicates.get(item.id) is not None:
            continue
        if item.group not in LANGS or not item.title.strip():
            continue
        split = split_of(edition_of(item.first_seen_at), dev, test)
        if split:
            out[item.id] = split
    return out


def _vectors(items: list[Item]) -> dict[int, dict[str, float]]:
    vectors, _ = _vectorize([tokenize(_document_text(i), i.group) for i in items])
    return {item.id: vec for item, vec in zip(items, vectors)}


def similar_pairs(items: list[Item], vectors: dict[int, dict[str, float]], min_similarity: float,
                  max_df_ratio: float = 0.05) -> dict[tuple[int, int], float]:
    """TF-IDF 相似度 ≥ 門檻、發布時間差 ≤ 48 小時的配對。候選只用少見 token 找（同聚合程式），再算精確分數。"""
    by_id = {i.id: i for i in items}
    df: dict[str, int] = defaultdict(int)
    for vec in vectors.values():
        for token in vec:
            df[token] += 1
    cap = max(20, int(len(items) * max_df_ratio))
    index: dict[str, list[int]] = defaultdict(list)
    for item_id, vec in vectors.items():
        for token in vec:
            if df[token] <= cap:
                index[token].append(item_id)
    out = {}
    for ids in index.values():
        for x in range(len(ids)):
            for y in range(x + 1, len(ids)):
                key = (min(ids[x], ids[y]), max(ids[x], ids[y]))
                if key in out:
                    continue
                a, b = by_id[key[0]], by_id[key[1]]
                if abs(a.published_at - b.published_at) > MAX_GAP:
                    continue
                score = _cosine(vectors[key[0]], vectors[key[1]])
                if score >= min_similarity:
                    out[key] = score
    return out


def build_pairs(items: list[Item], eligible: dict[int, str], comp_a: Components, comp_c: Components,
                min_similarity: float = MIN_SIMILARITY) -> list[Pair]:
    """所有候選配對（未分層抽樣前）：A 或 C 合併的，加上兩者都不合併但相似度夠的。"""
    pool = [i for i in items if i.id in eligible]
    by_id = {i.id: i for i in pool}
    vectors = _vectors(pool)
    keys: dict[tuple[int, int], float] = {}
    for comp in (comp_a, comp_c):
        members: dict[tuple, list[int]] = defaultdict(list)
        for item in pool:
            members[comp.root(item.id)].append(item.id)
        for ids in members.values():
            ids.sort()
            for x in range(len(ids)):
                for y in range(x + 1, len(ids)):
                    keys.setdefault((ids[x], ids[y]), -1.0)
    keys.update(similar_pairs(pool, vectors, min_similarity))
    pairs = []
    for (x, y), score in sorted(keys.items()):
        a, b = by_id[x], by_id[y]
        if eligible[x] != eligible[y] or abs(a.published_at - b.published_at) > MAX_GAP:
            continue
        if score < 0:
            score = _cosine(vectors[x], vectors[y])
        pairs.append(Pair(x, y, eligible[x], "same" if a.group == b.group else "cross",
                          comp_a.merged(x, y), comp_c.merged(x, y), round(score, 4)))
    for p in pairs:
        p.stratum = stratum_of(p, min_similarity)
    return pairs


def stratum_of(pair: Pair, min_similarity: float = MIN_SIMILARITY) -> str | None:
    if pair.a_merged and pair.c_merged:
        return "both"
    if pair.a_merged:
        return "only_a"
    if pair.c_merged:
        return "only_c"
    return "neither" if pair.similarity >= min_similarity else None


def fill_pairs(items: list[Item], eligible: dict[int, str], comp_a: Components, comp_c: Components,
               split: str, exclude: set[tuple[int, int]], rng: random.Random, count: int) -> list[Pair]:
    """中英某層不夠時的補抽：48 小時內、同版別（`section`）的隨機中英配對。"""
    pool = [i for i in items if eligible.get(i.id) == split]
    zh = [i for i in pool if i.group == "zh"]
    en = [i for i in pool if i.group == "en"]
    candidates = sorted(
        (min(a.id, b.id), max(a.id, b.id))
        for a in zh for b in en
        if a.section == b.section and abs(a.published_at - b.published_at) <= MAX_GAP
    )
    candidates = [k for k in candidates if k not in exclude]
    picked = rng.sample(candidates, min(count, len(candidates)))
    by_id = {i.id: i for i in pool}
    out = []
    for x, y in picked:
        vec = _vectors([by_id[x], by_id[y]])
        out.append(Pair(x, y, split, "cross", comp_a.merged(x, y), comp_c.merged(x, y),
                        round(_cosine(vec[x], vec[y]), 4), stratum=None, filled=True))
        out[-1].stratum = stratum_of(out[-1], 0.0)
    return out


# ── 錯誤案例 ─────────────────────────────────────────────────────────────


_SPACE = re.compile(r"\s+")


def norm_title(title: str | None) -> str:
    return _SPACE.sub("", zhconv.convert(unicodedata.normalize("NFKC", title or ""), "zh-tw"))


def error_case_title_pairs(cases_dir: str | Path) -> list[tuple[str, list[tuple[str, str]]]]:
    """回歸案例裡「應該／不應該是同一事件」的標題配對：每案一串（案例 id, [(標題, 標題)]）。"""
    out = []
    for path in sorted(Path(cases_dir).glob("E*.json")):
        case = json.loads(path.read_text(encoding="utf-8"))
        if not case.get("is_error"):
            continue
        pairs = []
        for event in case.get("event_titles", []):
            pairs += [(event, banned) for banned in case.get("must_not_contain_titles", [])]
        for group in case.get("must_merge_titles", []):
            members = list(group) + list(case.get("same_event_titles", []))
            pairs += [(members[i], members[j]) for i in range(len(members)) for j in range(i + 1, len(members))]
        pairs += [tuple(p) for p in case.get("reprint_pairs", [])]
        if pairs:
            out.append((case["case_id"], pairs))
    return out


def error_case_pairs(items: list[Item], title_pairs: list[tuple[str, list[tuple[str, str]]]]) -> tuple[
    dict[str, list[tuple[int, int]]], set[int]
]:
    """標題配對 → 素材 id 配對；另外回傳所有錯誤案例素材 id（test 一律排除）。"""
    by_title: dict[str, list[int]] = defaultdict(list)
    for item in items:
        by_title[norm_title(item.title)].append(item.id)
    pairs: dict[str, list[tuple[int, int]]] = {}
    involved: set[int] = set()
    for case_id, titles in title_pairs:
        found = []
        for t1, t2 in titles:
            for x in by_title.get(norm_title(t1), []):
                for y in by_title.get(norm_title(t2), []):
                    if x != y:
                        found.append((min(x, y), max(x, y)))
                        involved |= {x, y}
        pairs[case_id] = sorted(set(found))
    return pairs, involved


# ── 抽樣 ────────────────────────────────────────────────────────────────


def sample(pairs: list[Pair], *, rng: random.Random, quotas: dict = QUOTAS,
           forced: dict[str, list[tuple[int, int]]] | None = None, error_items: set[int] = frozenset(),
           filler=None, b_per_split: int = B_PER_SPLIT) -> SampleResult:
    """分層隨機抽樣。`forced`：每個錯誤案例的候選配對（只進 dev、每案最多 1 組）。
    `filler(split, exclude, count) -> list[Pair]`：中英層不足時補抽。"""
    by_key = {p.key: p for p in pairs}
    selected: list[Pair] = []
    shortages = []
    pools: dict[str, int] = {}
    taken: set[tuple[int, int]] = set()

    # 錯誤案例：每案隨機挑 1 組可標的 dev 配對，先占該層 dev 名額
    forced_by_stratum: dict[tuple[str, str], list[Pair]] = defaultdict(list)
    for case_id, keys in sorted((forced or {}).items()):
        options = [by_key[k] for k in keys if k in by_key and by_key[k].split == "dev" and by_key[k].stratum]
        if options:
            choice = rng.choice(options)
            choice.error_case = case_id
            forced_by_stratum[(choice.kind, choice.stratum)].append(choice)

    for split in ("dev", "test"):
        for kind in ("same", "cross"):
            for stratum in STRATA:
                quota = quotas[split][kind][stratum]
                pool = sorted(
                    (p for p in pairs if p.split == split and p.kind == kind and p.stratum == stratum
                     and p.key not in taken and not (split == "test" and {p.a, p.b} & error_items)),
                    key=lambda p: p.key,
                )
                pools[f"{split}/{kind}/{stratum}"] = len(pool)
                chosen = forced_by_stratum.get((kind, stratum), [])[:quota] if split == "dev" else []
                rest = [p for p in pool if p not in chosen and p.error_case is None]
                chosen = chosen + rng.sample(rest, min(quota - len(chosen), len(rest)))
                if len(chosen) < quota and kind == "cross" and filler is not None:
                    exclude = taken | {p.key for p in chosen}
                    extra = [p for p in filler(split, exclude, quota - len(chosen))
                             if not (split == "test" and {p.a, p.b} & error_items)]
                    chosen += extra
                if len(chosen) < quota:
                    shortages.append({"split": split, "kind": kind, "stratum": stratum, "quota": quota, "got": len(chosen)})
                for p in chosen:
                    taken.add(p.key)
                selected += chosen

    b_pairs = []
    for split in ("dev", "test"):
        members = sorted((p for p in selected if p.split == split), key=lambda p: p.key)
        b_pairs += rng.sample(members, min(b_per_split, len(members)))

    numbers = rng.sample(range(100, 1000), len(selected))
    pair_ids = {p.key: f"{'D' if p.split == 'dev' else 'T'}{n}" for p, n in zip(selected, numbers)}
    return SampleResult(selected, b_pairs, pair_ids, shortages, pools, {})


# ── 匯出 ────────────────────────────────────────────────────────────────


_TAG = re.compile(r"<[^>]+>")


def clean_summary(text: str, limit: int = SUMMARY_CHARS) -> str:
    text = _SPACE.sub(" ", html.unescape(_TAG.sub(" ", text or ""))).strip()
    return text if len(text) <= limit else text[: limit - 1] + "…"


def taipei_time(value: datetime) -> str:
    return value.astimezone(TAIPEI).strftime("%Y-%m-%d %H:%M")


def annotation_rows(pairs: list[Pair], items: dict[int, Item], pair_ids: dict, rng: random.Random) -> list[list[str]]:
    """給標註員的列：只有 pair_id、兩則的標題／摘要／台北時間、空白的判斷與理由。順序打亂、左右隨機。"""
    rows = []
    for split in ("dev", "test"):
        chunk = [p for p in pairs if p.split == split]
        chunk = rng.sample(chunk, len(chunk))
        for p in chunk:
            first, second = (p.a, p.b) if rng.random() < 0.5 else (p.b, p.a)
            x, y = items[first], items[second]
            rows.append([pair_ids[p.key], x.title, clean_summary(x.summary), taipei_time(x.published_at),
                         y.title, clean_summary(y.summary), taipei_time(y.published_at), "", ""])
    return rows


def write_csv(path: Path, rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(ANNOTATION_COLUMNS)
        writer.writerows(rows)


def write_xlsx(path: Path, sheets: dict[str, list[list[str]]], instructions: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    info = wb.active
    info.title = "說明"
    for line in instructions:
        info.append([line])
    info.column_dimensions["A"].width = 110
    widths = (8, 40, 60, 17, 40, 60, 17, 10, 30)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(list(ANNOTATION_COLUMNS))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        for col, width in zip("ABCDEFGHI", widths):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if rows:
            dv = DataValidation(type="list", formula1=f'"{",".join(JUDGMENTS)}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"H2:H{len(rows) + 1}")
        ws.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


INSTRUCTIONS = [
    "實驗一：兩則素材是不是「同一事件」（手冊第 9.3 節）。",
    "同一事件＝同一個主角的同一個具體行動或結果（48 小時內）；更新稿、轉載、以這件事為主體的分析稿也算「是」。",
    "同主題但不同行動、由這件事引起的後續事件、例行數據的不同日期、預告與結果相隔超過 48 小時：「不是」。",
    "摘要空白只有標題、或要靠背景知識才能判斷：「不確定」，並寫一句理由。",
    "「判斷」欄只填：是／不是／不確定。「一句理由」盡量填。不要查網路、不要用 LLM、不要和另一位標註員討論（第一次要獨立填）。",
    "pair_id 開頭 D＝dev（11/07 前交）、T＝test（11/12 前交）。時間是台北時間。",
]


def sha256_of(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()
