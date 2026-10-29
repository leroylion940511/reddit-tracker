"""實驗一抽樣（reddit_tracker/eval/exp1.py）：小型合成資料，不呼叫 LLM、不碰任何真實 DB。"""

from __future__ import annotations

import csv
import random
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest

from reddit_tracker.collection.entities import load_lexicon
from reddit_tracker.eval import exp1
from reddit_tracker.eval.exp1 import Components, Item, Pair

CASES_DIR = Path(__file__).parent / "regression" / "cases"
T0 = datetime(2026, 10, 27, 2, 0, tzinfo=timezone.utc)  # 10/27 10:00 台北 → 出刊日 10/28（dev）
T_TEST = datetime(2026, 10, 31, 2, 0, tzinfo=timezone.utc)  # 出刊日 11/01（test）


def _item(i, title, *, lang="en", seen=T0, source="bbc-world", section="world", cluster=None, dup=None, summary=""):
    return Item(i, source, section, "T2", title, summary, lang, f"https://x/{i}", f"https://x/{i}", seen, False,
                seen, dup, cluster)


# ── 日期切分 ─────────────────────────────────────────────────────────────


def test_edition_of_uses_0600_taipei_cutoff():
    assert exp1.edition_of(datetime(2026, 10, 26, 22, 0, tzinfo=timezone.utc)) == date(2026, 10, 27)
    assert exp1.edition_of(datetime(2026, 10, 26, 22, 0, 1, tzinfo=timezone.utc)) == date(2026, 10, 28)
    assert exp1.split_of(date(2026, 10, 27)) == "dev" and exp1.split_of(date(2026, 10, 30)) == "dev"
    assert exp1.split_of(date(2026, 10, 31)) == "test" and exp1.split_of(date(2026, 11, 2)) == "test"
    assert exp1.split_of(date(2026, 10, 26)) is None and exp1.split_of(date(2026, 11, 3)) is None


def test_eligible_items_drop_disabled_duplicates_other_languages_and_out_of_range():
    items = [
        _item(1, "A"), _item(2, "B", source="ptt"), _item(3, "C"), _item(4, "D", lang="ja"),
        _item(5, "", lang="en"), _item(6, "E", seen=T_TEST), _item(7, "F", seen=T0 - timedelta(days=3)),
    ]
    eligible = exp1.eligible_items(items, frozenset({"ptt"}), {3: 1}, exp1.DEV_DATES, exp1.TEST_DATES)
    assert eligible == {1: "dev", 6: "test"}


def test_components_are_transitive_across_merge_groups():
    comp = Components({1: 10, 2: 10, 3: 20, 4: 30, 5: None}, [[20, 10], [30, 40]])
    assert comp.merged(1, 2) and comp.merged(1, 3)  # 經跨語言合併組間接相連
    assert not comp.merged(1, 4)
    assert not comp.merged(5, 1) and comp.merged(5, 5)


def test_replay_groups_maps_live_clusters_through_items():
    live = {403: [1053, 1054], 410: [2000]}
    cluster_a = {1053: 7, 1054: 7, 2000: 9}
    assert exp1.replay_groups([[403, 410], [999, 410]], live, cluster_a) == [[7, 9]]


def test_method_a_is_the_old_threshold_without_entities():
    assert exp1.METHOD_A_RULES.required_entities(0.30) == 0
    assert exp1.METHOD_A_RULES.required_entities(0.90) == 0
    assert exp1.METHOD_A_RULES.required_entities(0.27) is None


def test_build_pairs_never_crosses_dev_and_test():
    items = [
        _item(1, "Serbia Vucic concedes defeat in parliamentary election", cluster=1),
        _item(2, "Serbia Vucic concedes defeat after parliamentary election", cluster=1),
        _item(3, "Serbia Vucic concedes defeat parliamentary election result", seen=T0 + timedelta(hours=40)),
    ]
    eligible = {1: "dev", 2: "dev", 3: "test"}
    comp = Components({1: 1, 2: 1, 3: 1}, [])
    pairs = exp1.build_pairs(items, eligible, comp, comp)
    assert [(p.a, p.b, p.split, p.stratum) for p in pairs] == [(1, 2, "dev", "both")]


def test_simulate_reproduces_old_and_new_rules_offline():
    """同一組素材：舊規則（A）只看相似度，新規則（C）還要共同實體。"""
    from reddit_tracker.collection.entities import MergeRules
    from reddit_tracker.config import get_settings
    from reddit_tracker.seeds.news_sources import load_registry

    registry = load_registry(get_settings().news_sources_path)
    items = [
        _item(1, "Microsoft offers buyouts to staff in cost cuts", source="ft-home"),
        _item(2, "Skydance offers buyouts to staff in cost cuts", source="reuters-business", seen=T0 + timedelta(minutes=5)),
    ]
    end = T0 + timedelta(hours=1)
    a, _ = exp1.simulate(items, rules=exp1.METHOD_A_RULES, lexicon=load_lexicon(), registry=registry,
                         recompute=False, disabled_sources=frozenset(), end=end, round_hours=1)
    c, _ = exp1.simulate(items, rules=MergeRules(), lexicon=load_lexicon(), registry=registry,
                         recompute=True, disabled_sources=frozenset(), end=end, round_hours=1)
    assert a[1] == a[2]  # 舊規則：buyouts 相似就併（E04 的錯）
    assert c[1] != c[2]  # 新規則：沒有共同實體不併


# ── 抽樣 ────────────────────────────────────────────────────────────────

SMALL = {
    "dev": {"same": dict.fromkeys(exp1.STRATA, 3), "cross": dict.fromkeys(exp1.STRATA, 2)},
    "test": {"same": dict.fromkeys(exp1.STRATA, 3), "cross": dict.fromkeys(exp1.STRATA, 2)},
}
FLAGS = {"both": (True, True, 0.6), "only_a": (True, False, 0.3), "only_c": (False, True, 0.3), "neither": (False, False, 0.2)}


def _pairs(n_same=6, n_cross=4, cross_short=None):
    pairs, k = [], 0
    for split in ("dev", "test"):
        for kind, n in (("same", n_same), ("cross", n_cross)):
            for stratum, (a, c, sim) in FLAGS.items():
                count = cross_short if (kind == "cross" and stratum == "neither" and cross_short is not None) else n
                for _ in range(count):
                    k += 2
                    p = Pair(k, k + 1, split, kind, a, c, sim)
                    p.stratum = exp1.stratum_of(p)
                    pairs.append(p)
    return pairs


def test_stratified_counts_b_double_and_pair_ids():
    result = exp1.sample(_pairs(), rng=random.Random(1), quotas=SMALL, b_per_split=5)
    for split in ("dev", "test"):
        for kind, quota in (("same", 3), ("cross", 2)):
            for stratum in exp1.STRATA:
                got = [p for p in result.selected if (p.split, p.kind, p.stratum) == (split, kind, stratum)]
                assert len(got) == quota
    assert len(result.selected) == 40 and not result.shortages
    assert [p.split for p in result.b_pairs].count("dev") == 5 and [p.split for p in result.b_pairs].count("test") == 5
    assert {p.key for p in result.b_pairs} <= {p.key for p in result.selected}
    ids = list(result.pair_ids.values())
    assert len(set(ids)) == len(ids)
    assert all(result.pair_ids[p.key][0] == ("D" if p.split == "dev" else "T") for p in result.selected)
    # pair_id 的數字是隨機的，跟分層無關
    numbers = [int(result.pair_ids[p.key][1:]) for p in result.selected]
    assert numbers != sorted(numbers)


def test_same_seed_same_sample_different_seed_differs():
    def keys(seed):
        r = exp1.sample(_pairs(), rng=random.Random(seed), quotas=SMALL, b_per_split=5)
        return [(p.key, r.pair_ids[p.key]) for p in r.selected], [p.key for p in r.b_pairs]

    assert keys(exp1.SEED) == keys(exp1.SEED)
    assert keys(exp1.SEED) != keys(exp1.SEED + 1)


def test_cross_language_shortage_is_filled_but_same_language_is_not():
    calls = []

    def filler(split, exclude, count):
        calls.append((split, count))
        out = []
        for i in range(count):
            p = Pair(10_000 + len(calls) * 10 + i, 20_000 + i, split, "cross", False, False, 0.0, filled=True)
            p.stratum = exp1.stratum_of(p, 0.0)
            out.append(p)
        return out

    pairs = [p for p in _pairs(cross_short=1) if not (p.kind == "same" and p.stratum == "only_c" and p.split == "dev")]
    result = exp1.sample(pairs, rng=random.Random(3), quotas=SMALL, filler=filler, b_per_split=5)
    assert calls == [("dev", 1), ("test", 1)]  # 中英「都不合併」層各缺 1 組
    filled = [p for p in result.selected if p.filled]
    assert len(filled) == 2 and all(p.kind == "cross" for p in filled)
    assert result.shortages == [{"split": "dev", "kind": "same", "stratum": "only_c", "quota": 3, "got": 0}]


def test_error_cases_go_to_dev_only():
    pairs = _pairs()
    dev_both = next(p for p in pairs if (p.split, p.kind, p.stratum) == ("dev", "same", "only_a"))
    test_pair = next(p for p in pairs if p.split == "test")
    forced = {"E01": [dev_both.key], "E99": [test_pair.key]}
    error_items = {dev_both.a, dev_both.b, test_pair.a}
    result = exp1.sample(pairs, rng=random.Random(5), quotas=SMALL, forced=forced, error_items=error_items, b_per_split=5)
    chosen = {p.key: p for p in result.selected}
    assert chosen[dev_both.key].error_case == "E01"
    assert test_pair.key not in chosen  # test 不放錯誤案例素材
    assert all(not ({p.a, p.b} & error_items) for p in result.selected if p.split == "test")
    assert {p.error_case for p in result.selected if p.error_case} == {"E01"}


def test_error_case_title_pairs_come_from_regression_cases():
    cases = dict(exp1.error_case_title_pairs(CASES_DIR))
    assert any("Seahawks" in t2 for _, t2 in cases["E01"])  # 塞爾維亞事件 × AP 美式足球稿
    assert "E09" not in cases  # 不是錯誤案例
    assert {"E01", "E02", "E03", "E04", "E05", "E06", "E07", "E10"} <= set(cases)
    items = [_item(1, cases["E04"][0][0]), _item(2, cases["E04"][0][1])]
    pairs, involved = exp1.error_case_pairs(items, [("E04", cases["E04"][:1])])
    assert pairs == {"E04": [(1, 2)]} and involved == {1, 2}


# ── 匯出：沒有洩漏欄位 ───────────────────────────────────────────────────


def test_annotation_export_has_no_leaking_columns(tmp_path):
    items = {
        1: _item(1, "Vucic concedes", summary="<p>Serbia&amp;co</p>"),
        2: _item(2, "武契奇承認敗選", lang="zh-Hant", summary="塞爾維亞" * 200, seen=T0 + timedelta(hours=1)),
    }
    pair = Pair(1, 2, "dev", "cross", True, False, 0.4321, stratum="only_a")
    rows = exp1.annotation_rows([pair], items, {pair.key: "D417"}, random.Random(0))
    assert len(rows) == 1 and rows[0][0] == "D417" and rows[0][7:] == ["", ""]
    flat = " ".join(rows[0])
    for leak in ("only_a", "0.4321", "both", "neither", "A、C", "stratum", "merged"):
        assert leak not in flat
    assert "2026-10-27 10:00" in rows[0] and "Serbia&co" in rows[0]
    long = next(c for c in rows[0] if c.startswith("塞爾維亞"))
    assert len(long) == exp1.SUMMARY_CHARS and long.endswith("…")

    exp1.write_csv(tmp_path / "a.csv", rows)
    with (tmp_path / "a.csv").open(encoding="utf-8-sig") as fh:
        assert next(csv.reader(fh)) == list(exp1.ANNOTATION_COLUMNS)
    exp1.write_xlsx(tmp_path / "a.xlsx", {"dev": rows, "test": []}, exp1.INSTRUCTIONS)
    from openpyxl import load_workbook

    wb = load_workbook(tmp_path / "a.xlsx")
    assert wb.sheetnames == ["說明", "dev", "test"]
    header = [c.value for c in wb["dev"][1]]
    assert header == list(exp1.ANNOTATION_COLUMNS)
    assert wb["dev"]["A2"].value == "D417"


def test_annotation_rows_are_shuffled_reproducibly():
    items = {i: _item(i, f"title {i}") for i in range(1, 41)}
    pairs = [Pair(i, i + 1, "dev", "same", False, False, 0.2, stratum="neither") for i in range(1, 40, 2)]
    ids = {p.key: f"D{100 + n}" for n, p in enumerate(pairs)}
    first = exp1.annotation_rows(pairs, items, ids, random.Random(9))
    again = exp1.annotation_rows(pairs, items, ids, random.Random(9))
    assert first == again
    assert [r[0] for r in first] != [ids[p.key] for p in pairs]


@pytest.mark.parametrize("seed", [exp1.SEED])
def test_seed_is_committed_constant(seed):
    assert seed == 20261104
